# -*- coding: utf-8 -*-
"""Portal Inmobiliario (Portalinmobiliario.com) scraper

- Playwright browser automation
- Parses Nordic JSON from __NORDIC_RENDERING_CTX__ / _n.ctx.r
- Paginates using pagination.next_page.url (pattern _Desde_49_ ...)
- Exports formatted Excel (no #######)
- Avoids PermissionError by saving to a unique filename when locked
- Checkpointing every N rows

Usage:
  pip install -r requirements.txt
  python searchdatahouse.py

Edit MAX_LISTINGS to control how many listings to collect.
"""

import json
import os
import re
import sys
import time
import unicodedata
from urllib.error import URLError
from urllib.parse import parse_qs, unquote_plus, urlparse
from urllib.request import urlopen
from datetime import datetime, date
from pathlib import Path

import pandas as pd

from playwright.sync_api import sync_playwright

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIG
# ============================================================
DEFAULT_LISTING_URL = (
    "https://www.portalinmobiliario.com/venta/casa/propiedades-usadas/"
    "las-condes-metropolitana/_OrderId_PRICE_PriceRange_0CLP-400000000CLP_NoIndex_True"
)
LISTING_URL = os.environ.get("LISTING_URL", DEFAULT_LISTING_URL)

CITY = "las_condes"
PROPERTY_TYPE = os.environ.get("PROPERTY_TYPE", "auto").lower()
MAX_LISTINGS = int(os.environ.get("MAX_LISTINGS", "300"))
# Headless por defecto: Chrome corre sin ventana visible.
HEADLESS = os.environ.get("HEADLESS", "true").lower() == "true"
BROWSER_CHANNEL = os.environ.get("BROWSER_CHANNEL", "chrome")
CHECKPOINT_EVERY_N = int(os.environ.get("CHECKPOINT_EVERY_N", "50"))
PRICE_PER_M2_THRESHOLD_UF = 60
UF_REFRESH_MODE = "per_run"
UF_API_URL = "https://mindicador.cl/api/uf"
BLOCK_HEAVY_RESOURCES = os.environ.get("BLOCK_HEAVY_RESOURCES", "true").lower() == "true"
INITIAL_PAGE_DELAY_SECONDS = float(os.environ.get("INITIAL_PAGE_DELAY_SECONDS", "0.2"))
AFTER_NORDIC_DELAY_SECONDS = float(os.environ.get("AFTER_NORDIC_DELAY_SECONDS", "0.1"))
NEXT_PAGE_DELAY_SECONDS = float(os.environ.get("NEXT_PAGE_DELAY_SECONDS", "0.5"))
ROW_DELAY_SECONDS = float(os.environ.get("ROW_DELAY_SECONDS", "0"))
OVERLAY_DISMISS_TRIES = int(os.environ.get("OVERLAY_DISMISS_TRIES", "2"))
SAVE_DEBUG_HTML = os.environ.get("SAVE_DEBUG_HTML", "false").lower() == "true"

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/135.0 Safari/537.36"
)

DEBUG_HTML_FILE = os.environ.get("DEBUG_HTML_FILE", "debug_playwright.html")
CHECKPOINT_FILE_ENV = os.environ.get("CHECKPOINT_FILE")
OUTPUT_XLSX_ENV = os.environ.get("OUTPUT_XLSX")
CHECKPOINT_FILE = os.environ.get(
    "CHECKPOINT_FILE",
    f"checkpoint_{CITY}_{date.today().isoformat()}.json",
)
OUTPUT_XLSX = os.environ.get(
    "OUTPUT_XLSX",
    f"properties_{CITY}_{date.today().isoformat()}.xlsx",
)


# ============================================================
# UTILS
# ============================================================
def iso_now():
    return datetime.now().isoformat(timespec="seconds")


def fixed_sleep(seconds):
    if seconds > 0:
        time.sleep(seconds)


def clean_number(text):
    if not text:
        return None
    digits = re.sub(r"[^\d]", "", str(text))
    return int(digits) if digits else None


def as_number(value):
    if isinstance(value, (int, float)):
        return value
    return None


def normalize_url(u):
    if not u:
        return ""
    if u.startswith("http"):
        return u
    return "https://www." + u.lstrip("/")


def slugify_text(text, default="busqueda", max_length=90):
    normalized = unicodedata.normalize("NFKD", str(text or ""))
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", ascii_text.lower()).strip("_")
    slug = re.sub(r"_+", "_", slug)
    return (slug[:max_length].strip("_") or default)


def get_path_segment_after_property_type(url, property_type):
    parts = [unquote_plus(p) for p in urlparse(url).path.split("/") if p]
    if property_type in parts:
        idx = parts.index(property_type)
        if idx + 1 < len(parts):
            return parts[idx + 1]
    return CITY


def get_applied_filter_parts(url):
    fragment = unquote_plus(urlparse(url).fragment or "")
    params = parse_qs(fragment)
    filter_name = (params.get("applied_filter_name") or [""])[0]
    value_name = (params.get("applied_value_name") or [""])[0]

    parts = []
    if filter_name:
        parts.append(filter_name)
    if value_name:
        parts.append(value_name)
    return parts


def build_search_slug(url):
    property_type = detect_property_type(url)
    if property_type == "auto":
        property_type = "propiedades"

    location = get_path_segment_after_property_type(url, property_type)
    location = re.sub(r"-metropolitana$", "", location)

    parts = ["properties", property_type, location]
    parts.extend(get_applied_filter_parts(url))
    return slugify_text("_".join(parts), default=f"properties_{CITY}")


def configure_run_filenames(url):
    global CHECKPOINT_FILE, OUTPUT_XLSX

    search_slug = build_search_slug(url)
    today = date.today().isoformat()

    if not CHECKPOINT_FILE_ENV:
        CHECKPOINT_FILE = f"checkpoint_{search_slug}_{today}.json"
    if not OUTPUT_XLSX_ENV:
        OUTPUT_XLSX = f"{search_slug}_{today}.xlsx"


def choose_listing_url(default_url):
    if "LISTING_URL" in os.environ or not sys.stdin.isatty():
        return default_url

    print("\n=== Portal Inmobiliario Scraper ===")
    print("1) Usar link por defecto")
    print("2) Ingresar otro link del Portal Inmobiliario")

    try:
        option = input("Elige una opcion [1]: ").strip()
    except EOFError:
        return default_url

    if option in ("", "1"):
        return default_url

    if option == "2":
        try:
            url = input("Pega el link del Portal Inmobiliario: ").strip()
        except EOFError:
            return default_url

        if url.startswith("http"):
            return url

        print("Link no valido. Se usara el link por defecto.")
        return default_url

    print("Opcion no valida. Se usara el link por defecto.")
    return default_url


def detect_block_or_captcha(html: str) -> bool:
    if not html:
        return False
    s = html.lower()
    keywords = [
        "captcha", "verifica que eres", "verify you are", "robot",
        "soy humano", "challenge", "security check"
    ]
    return any(k in s for k in keywords)


def detect_property_type(url):
    low = (url or "").lower()
    if "/departamento/" in low:
        return "departamento"
    if "/casa/" in low:
        return "casa"
    return "auto"


def matches_property_type(headline, title, url, property_type):
    if property_type == "auto":
        return True

    haystack = " ".join([headline or "", title or "", url or ""]).lower()
    if property_type == "departamento":
        return "departamento" in haystack or "/departamento/" in haystack
    if property_type == "casa":
        return "casa" in haystack or "/casa/" in haystack
    return True


def empty_checkpoint_state():
    return {
        "listing_url": LISTING_URL,
        "page_number": 1,
        "current_url": LISTING_URL,
        "completed_rows": 0,
        "scraped_urls": [],
        "rows": [],
        "errors": []
    }


def load_checkpoint():
    p = Path(CHECKPOINT_FILE)
    if not p.exists():
        return empty_checkpoint_state()
    try:
        state = json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return empty_checkpoint_state()

    checkpoint_url = state.get("listing_url")
    if checkpoint_url and checkpoint_url != LISTING_URL:
        print("Checkpoint ignorado: pertenece a otro link de busqueda.")
        return empty_checkpoint_state()
    if not checkpoint_url and LISTING_URL != DEFAULT_LISTING_URL:
        print("Checkpoint antiguo ignorado: no identifica el link de busqueda.")
        return empty_checkpoint_state()

    state["listing_url"] = LISTING_URL
    return state


def save_checkpoint(state):
    Path(CHECKPOINT_FILE).write_text(
        json.dumps(state, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )


def fetch_current_uf_clp():
    try:
        with urlopen(UF_API_URL, timeout=15) as response:
            data = json.loads(response.read().decode("utf-8"))
    except (OSError, URLError, json.JSONDecodeError) as exc:
        raise RuntimeError(f"No se pudo obtener la UF desde {UF_API_URL}: {exc}") from exc

    serie = data.get("serie") or []
    if not serie:
        raise RuntimeError(f"La respuesta de UF desde {UF_API_URL} no trae datos en 'serie'.")

    value = serie[0].get("valor")
    if not isinstance(value, (int, float)):
        raise RuntimeError(f"La respuesta de UF desde {UF_API_URL} no trae un valor numerico valido.")

    return float(value)


def calculate_price_metrics(price_clp, price_uf, useful_area, uf_value_clp):
    normalized_clp = as_number(price_clp)
    normalized_uf = as_number(price_uf)

    if normalized_clp is None and as_number(price_uf) is not None:
        if uf_value_clp is None:
            return None, normalized_uf, None, None
        normalized_clp = round(price_uf * uf_value_clp)

    if normalized_uf is None and normalized_clp is not None and uf_value_clp:
        normalized_uf = normalized_clp / uf_value_clp

    area = as_number(useful_area)
    if normalized_clp is None and normalized_uf is None:
        return None, None, None, None
    if area is None or area <= 0:
        return normalized_clp, normalized_uf, None, None

    price_clp_per_m2 = round(normalized_clp / area) if normalized_clp is not None else None
    price_uf_per_m2 = round(normalized_uf / area, 2) if normalized_uf is not None else None
    return normalized_clp, normalized_uf, price_clp_per_m2, price_uf_per_m2


def add_price_metrics(row, uf_value_clp):
    price_normalized_clp, price_normalized_uf, price_clp_per_m2, price_uf_per_m2 = calculate_price_metrics(
        as_number(row.get("Price_CLP")),
        as_number(row.get("Price_UF")),
        as_number(row.get("Useful_Area_m2")),
        uf_value_clp,
    )
    row["UF_Value_CLP"] = uf_value_clp if uf_value_clp is not None else "Not listed"
    row["Price_Normalized_CLP"] = (
        price_normalized_clp if price_normalized_clp is not None else "Not listed"
    )
    row["Price_Normalized_UF"] = (
        price_normalized_uf if price_normalized_uf is not None else "Not listed"
    )
    row["Price_CLP_per_m2"] = (
        price_clp_per_m2 if price_clp_per_m2 is not None else "Not listed"
    )
    row["Price_UF_per_m2"] = (
        price_uf_per_m2 if price_uf_per_m2 is not None else "Not listed"
    )
    return row


def rows_need_uf(rows):
    return any(as_number(row.get("Price_UF")) is not None for row in rows)


# ============================================================
# FILE LOCK SAFE OUTPUT (fix PermissionError)
# ============================================================
def unique_name(path: str) -> str:
    p = Path(path)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return str(p.with_name(f"{p.stem}_{stamp}{p.suffix}"))


def safe_excel_path(path: str) -> str:
    p = Path(path)
    if not p.exists():
        return str(p)

    return unique_name(path)


# ============================================================
# OVERLAY DISMISS
# ============================================================
def dismiss_overlays(page, tries=OVERLAY_DISMISS_TRIES):
    xpaths = [
        "//button[contains(., 'Entendido')]",
        "//button[contains(., 'Aceptar')]",
        "//button[contains(., 'Acepto')]",
        "//button[contains(., 'Aceptar todo')]",
        "//button[contains(., 'Aceptar todas')]",
        "//button[contains(., 'Accept')]",
        "//button[contains(., 'Accept all')]",
        "//button[contains(., 'OK')]",
    ]
    for _ in range(tries):
        clicked = False
        for xp in xpaths:
            try:
                btns = page.locator(f"xpath={xp}")
                for idx in range(btns.count()):
                    b = btns.nth(idx)
                    try:
                        if b.is_visible() and b.is_enabled():
                            b.click(timeout=1000)
                            clicked = True
                            fixed_sleep(0.1)
                    except Exception:
                        pass
            except Exception:
                pass
        if not clicked:
            break


# ============================================================
# NORDIC JSON
# ============================================================
def wait_for_nordic_ctx(page, timeout=90):
    timeout_ms = timeout * 1000
    page.wait_for_selector("#__NORDIC_RENDERING_CTX__", state="attached", timeout=timeout_ms)
    page.wait_for_function(
        """() => {
            const el = document.querySelector("#__NORDIC_RENDERING_CTX__");
            return Boolean(el && el.textContent && el.textContent.includes("_n.ctx.r="));
        }""",
        timeout=timeout_ms,
    )


def extract_nordic_json_from_dom(page):
    txt = page.locator("#__NORDIC_RENDERING_CTX__").text_content(timeout=5000) or ""

    marker = "_n.ctx.r="
    idx = txt.find(marker)
    if idx == -1:
        return None

    payload = txt[idx + len(marker):].lstrip()
    decoder = json.JSONDecoder()
    obj, end = decoder.raw_decode(payload)
    return obj


def block_heavy_resource(route):
    resource_type = route.request.resource_type
    if resource_type in {"image", "media", "font"}:
        route.abort()
        return
    route.continue_()


def parse_polycards(nordic_data):
    results = (
        nordic_data.get("appProps", {})
                  .get("pageProps", {})
                  .get("initialState", {})
                  .get("results", [])
    )
    return [x for x in results if x.get("id") == "POLYCARD"]


def components_to_dict(components):
    out = {}
    for c in components or []:
        t = c.get("type")
        if t:
            out[t] = c.get(t, {})
    return out


def get_next_page_url(nordic_data):
    return (
        nordic_data.get("appProps", {})
                  .get("pageProps", {})
                  .get("initialState", {})
                  .get("pagination", {})
                  .get("next_page", {})
                  .get("url")
    )


# ============================================================
# EXCEL FORMATTING
# ============================================================
def format_excel(file_path: str, sheet_name="Listings"):
    wb = openpyxl.load_workbook(file_path)
    if sheet_name not in wb.sheetnames:
        wb.save(file_path)
        return

    ws = wb[sheet_name]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="D9E1F2")

    for c in range(1, ws.max_column + 1):
        cell = ws.cell(1, c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}

    wrap_cols = {"Title", "Address", "Seller", "Description", "Source_URL"}
    for name in wrap_cols:
        if name in headers:
            col = headers[name]
            for r in range(2, ws.max_row + 1):
                ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")

    clp_columns = {"Price_CLP", "UF_Value_CLP", "Price_Normalized_CLP", "Price_CLP_per_m2"}
    for name in clp_columns:
        if name in headers:
            col = headers[name]
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(r, col)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '"$"#,##0'

    if "Price_UF" in headers:
        col = headers["Price_UF"]
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, col)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '"UF"#,##0'

    uf_decimal_columns = {"Price_Normalized_UF", "Price_UF_per_m2"}
    for name in uf_decimal_columns:
        if name in headers:
            col = headers[name]
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(r, col)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '"UF"#,##0.00'

    if "Price_USD" in headers:
        col = headers["Price_USD"]
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, col)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '"US$"#,##0.00'

    alert_fill = PatternFill("solid", fgColor="FFC7CE")
    alert_font = Font(color="9C0006")
    alert_cells = set()
    if "Price_UF_per_m2" in headers:
        col = headers["Price_UF_per_m2"]
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, col)
            if isinstance(cell.value, (int, float)) and cell.value < PRICE_PER_M2_THRESHOLD_UF:
                cell.fill = alert_fill
                cell.font = alert_font
                alert_cells.add((r, col))

    alt_fill = PatternFill("solid", fgColor="F7F7F7")
    for r in range(2, ws.max_row + 1):
        if r % 2 == 0:
            for c in range(1, ws.max_column + 1):
                if (r, c) in alert_cells:
                    continue
                ws.cell(r, c).fill = alt_fill

    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            v = ws.cell(row_idx, col_idx).value
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

    wb.save(file_path)


# ============================================================
# MAIN
# ============================================================
def main():
    global LISTING_URL, PROPERTY_TYPE
    LISTING_URL = choose_listing_url(LISTING_URL)
    configure_run_filenames(LISTING_URL)
    if PROPERTY_TYPE == "auto":
        PROPERTY_TYPE = detect_property_type(LISTING_URL)
    if PROPERTY_TYPE not in {"auto", "casa", "departamento"}:
        print(f"PROPERTY_TYPE={PROPERTY_TYPE} no es valido. Se usara auto.")
        PROPERTY_TYPE = "auto"

    state = load_checkpoint()
    uf_value_clp = None

    page_number = state.get("page_number", 1)
    current_url = state.get("current_url", LISTING_URL) or LISTING_URL

    scraped_urls = set(state.get("scraped_urls", []))
    rows = state.get("rows", [])
    errors = state.get("errors", [])

    duplicates_skipped = 0
    debug_html_written = False

    if UF_REFRESH_MODE == "per_run":
        try:
            uf_value_clp = fetch_current_uf_clp()
            print(f"UF actual obtenida: ${uf_value_clp:,.2f} CLP")
        except RuntimeError as exc:
            if rows_need_uf(rows):
                raise SystemExit(f"{exc}\nNo se pueden convertir precios en UF guardados en checkpoint.")
            print(f"Advertencia: {exc}")

    playwright = sync_playwright().start()
    browser = playwright.chromium.launch(
        channel=BROWSER_CHANNEL,
        headless=HEADLESS,
        args=[
            "--disable-blink-features=AutomationControlled",
            "--disable-dev-shm-usage",
            "--no-sandbox",
        ],
    )
    context = browser.new_context(
        user_agent=USER_AGENT,
        locale="es-CL",
        timezone_id="America/Santiago",
        viewport={"width": 1920, "height": 1080},
    )
    context.add_init_script(
        "Object.defineProperty(navigator, 'webdriver', {get: () => undefined});"
    )
    if BLOCK_HEAVY_RESOURCES:
        context.route("**/*", block_heavy_resource)
    page = context.new_page()
    page.set_default_timeout(30000)
    page.set_default_navigation_timeout(90000)

    try:
        while len(rows) < MAX_LISTINGS and current_url:
            page.goto(current_url, wait_until="domcontentloaded", timeout=90000)
            fixed_sleep(INITIAL_PAGE_DELAY_SECONDS)
            dismiss_overlays(page)

            wait_for_nordic_ctx(page, timeout=90)
            fixed_sleep(AFTER_NORDIC_DELAY_SECONDS)
            dismiss_overlays(page)

            html = page.content()
            if SAVE_DEBUG_HTML:
                Path(DEBUG_HTML_FILE).write_text(html, encoding="utf-8")
                debug_html_written = True

            if detect_block_or_captcha(html):
                Path(DEBUG_HTML_FILE).write_text(html, encoding="utf-8")
                debug_html_written = True
                page.screenshot(path="captcha_or_block.png", full_page=True)
                raise SystemExit("Bloqueo/CAPTCHA detectado. Screenshot: captcha_or_block.png")

            nordic_data = extract_nordic_json_from_dom(page)
            if not nordic_data:
                raise SystemExit(f"Nordic JSON no encontrado. Revisa {DEBUG_HTML_FILE}")

            polycards = parse_polycards(nordic_data)

            for item in polycards:
                if len(rows) >= MAX_LISTINGS:
                    break

                poly = item.get("polycard", {})
                meta = poly.get("metadata", {})
                comps = components_to_dict(poly.get("components", []))

                url = normalize_url(meta.get("url", ""))

                if url and url in scraped_urls:
                    duplicates_skipped += 1
                    continue

                headline = (comps.get("headline", {}).get("text") or "Not listed").strip()
                title = (comps.get("title", {}).get("text") or "Not listed").strip()
                address = (comps.get("location", {}).get("text") or "Not listed").strip()
                seller = (comps.get("seller", {}).get("text") or "Not listed").strip()

                title = re.sub(r"\s+", " ", title).strip()
                seller = re.sub(r"\s+", " ", seller).strip()

                if not matches_property_type(headline, title, url, PROPERTY_TYPE):
                    continue

                price_block = comps.get("price", {}).get("current_price", {}) or {}
                price_value = price_block.get("value")
                price_currency = price_block.get("currency")

                price_clp = price_value if price_currency == "CLP" else None
                price_uf = price_value if price_currency == "CLF" else None
                price_usd = price_value if price_currency == "USD" else None

                if price_uf is not None and uf_value_clp is None:
                    raise SystemExit(
                        "No se pudo convertir un precio en UF porque no hay valor UF disponible."
                    )

                attrs = comps.get("attributes_list", {}).get("texts", []) or []
                bedrooms = bathrooms = useful_area = None
                for t in attrs:
                    low = t.lower()
                    if "dorm" in low:
                        bedrooms = clean_number(low)
                    elif "bañ" in low:
                        bathrooms = clean_number(low)
                    elif "m²" in low or "m2" in low:
                        useful_area = clean_number(low)

                normalized_clp, normalized_uf, price_clp_per_m2, price_uf_per_m2 = calculate_price_metrics(
                    price_clp, price_uf, useful_area, uf_value_clp
                )

                rows.append({
                    "Headline": headline,
                    "Title": title,
                    "Address": address,
                    "Comuna": "Las Condes",
                    "Price_CLP": price_clp,
                    "Price_UF": price_uf,
                    "Price_USD": price_usd,
                    "UF_Value_CLP": uf_value_clp if uf_value_clp is not None else "Not listed",
                    "Price_Normalized_CLP": normalized_clp if normalized_clp is not None else "Not listed",
                    "Price_Normalized_UF": normalized_uf if normalized_uf is not None else "Not listed",
                    "Price_CLP_per_m2": price_clp_per_m2 if price_clp_per_m2 is not None else "Not listed",
                    "Price_UF_per_m2": price_uf_per_m2 if price_uf_per_m2 is not None else "Not listed",
                    "Bedrooms": bedrooms if bedrooms is not None else "Not listed",
                    "Bathrooms": bathrooms if bathrooms is not None else "Not listed",
                    "Useful_Area_m2": useful_area if useful_area is not None else "Not listed",
                    "Seller": seller,
                    "Description": "Not listed",
                    "Source_URL": url or "Not listed",
                    "Scraped_At": iso_now(),
                })

                if url:
                    scraped_urls.add(url)

                if len(rows) % CHECKPOINT_EVERY_N == 0:
                    save_checkpoint({
                        "listing_url": LISTING_URL,
                        "page_number": page_number,
                        "current_url": current_url,
                        "completed_rows": len(rows),
                        "scraped_urls": sorted(scraped_urls),
                        "rows": rows,
                        "errors": errors
                    })

                fixed_sleep(ROW_DELAY_SECONDS)

            next_url = get_next_page_url(nordic_data)
            if next_url:
                current_url = next_url
                page_number += 1
                fixed_sleep(NEXT_PAGE_DELAY_SECONDS)
            else:
                current_url = None

            save_checkpoint({
                "listing_url": LISTING_URL,
                "page_number": page_number,
                "current_url": current_url if current_url else "",
                "completed_rows": len(rows),
                "scraped_urls": sorted(scraped_urls),
                "rows": rows,
                "errors": errors
            })

    finally:
        context.close()
        browser.close()
        playwright.stop()

    if rows_need_uf(rows) and uf_value_clp is None:
        raise SystemExit("No se pueden exportar precios en UF porque no hay valor UF disponible.")
    rows = [add_price_metrics(row, uf_value_clp) for row in rows]

    df = pd.DataFrame(rows)
    err_df = pd.DataFrame(errors)

    columns = [
        "Headline", "Title", "Address", "Comuna",
        "Price_CLP", "Price_UF", "Price_USD",
        "UF_Value_CLP", "Price_Normalized_UF",
        "Price_UF_per_m2",
        "Bedrooms", "Bathrooms", "Useful_Area_m2",
        "Seller", "Source_URL", "Scraped_At"
    ]

    for c in columns:
        if c not in df.columns:
            df[c] = "Not listed"
    df = df[columns]

    out_path = safe_excel_path(OUTPUT_XLSX)

    try:
        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Listings")
            if not err_df.empty:
                err_df.to_excel(writer, index=False, sheet_name="Errors")
            else:
                pd.DataFrame([{
                    "Timestamp": "",
                    "Listing URL": "",
                    "Error Type": "",
                    "Error Message": "",
                    "Retry Count": ""
                }]).to_excel(writer, index=False, sheet_name="Errors")

            pd.DataFrame([{
                "Listing URL": LISTING_URL,
                "City": CITY,
                "Property Type": PROPERTY_TYPE,
                "Max Listings": MAX_LISTINGS,
                "Headless": HEADLESS,
                "Scraped At": iso_now(),
                "Debug HTML": DEBUG_HTML_FILE,
                "Checkpoint": CHECKPOINT_FILE,
                "UF Refresh Mode": UF_REFRESH_MODE,
                "UF Value CLP": uf_value_clp if uf_value_clp is not None else "Not listed",
                "Price per m2 Threshold UF": PRICE_PER_M2_THRESHOLD_UF,
                "Checkpoint Every N": CHECKPOINT_EVERY_N,
                "Block Heavy Resources": BLOCK_HEAVY_RESOURCES,
                "Row Delay Seconds": ROW_DELAY_SECONDS,
                "Next Page Delay Seconds": NEXT_PAGE_DELAY_SECONDS,
                "Save Debug HTML": SAVE_DEBUG_HTML,
                "Last Page Number": page_number,
                "Collected Rows": len(df),
                "Duplicates Skipped": duplicates_skipped
            }]).to_excel(writer, index=False, sheet_name="Session Info")

        format_excel(out_path, sheet_name="Listings")

    except PermissionError:
        alt = unique_name(OUTPUT_XLSX)
        with pd.ExcelWriter(alt, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Listings")
            if not err_df.empty:
                err_df.to_excel(writer, index=False, sheet_name="Errors")
            pd.DataFrame([{
                "Listing URL": LISTING_URL,
                "City": CITY,
                "Property Type": PROPERTY_TYPE,
                "Max Listings": MAX_LISTINGS,
                "Headless": HEADLESS,
                "Scraped At": iso_now(),
                "Debug HTML": DEBUG_HTML_FILE,
                "Checkpoint": CHECKPOINT_FILE,
                "UF Refresh Mode": UF_REFRESH_MODE,
                "UF Value CLP": uf_value_clp if uf_value_clp is not None else "Not listed",
                "Price per m2 Threshold UF": PRICE_PER_M2_THRESHOLD_UF,
                "Checkpoint Every N": CHECKPOINT_EVERY_N,
                "Block Heavy Resources": BLOCK_HEAVY_RESOURCES,
                "Row Delay Seconds": ROW_DELAY_SECONDS,
                "Next Page Delay Seconds": NEXT_PAGE_DELAY_SECONDS,
                "Save Debug HTML": SAVE_DEBUG_HTML,
                "Last Page Number": page_number,
                "Collected Rows": len(df),
                "Duplicates Skipped": duplicates_skipped
            }]).to_excel(writer, index=False, sheet_name="Session Info")

        format_excel(alt, sheet_name="Listings")
        out_path = alt

    print(f"✓ {len(df)} properties collected (unique)")
    print(f"✗ {len(errors)} errors (see Errors tab)")
    print(f"⏭ {duplicates_skipped} duplicates skipped")
    print(f"File ready: {out_path}")
    print(f"Checkpoint saved: {CHECKPOINT_FILE}")
    if debug_html_written:
        print(f"Debug HTML saved: {DEBUG_HTML_FILE}")
    else:
        print(f"Debug HTML disabled: set SAVE_DEBUG_HTML=true to write {DEBUG_HTML_FILE}")


if __name__ == "__main__":
    main()
